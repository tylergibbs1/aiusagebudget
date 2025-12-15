#!/usr/bin/env python3
"""
AI Usage Budget Tracker
Tracks Claude AI usage across multiple machines using bunx ccusage.
Sheets: Daily, Weekly, Monthly, Machine Totals
"""

import subprocess
import json
import hashlib
import platform
import re
import sys
import uuid

from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Constants
EXCEL_FILE = "usage_tracker.xlsx"

# Column headers for time-based sheets
TIME_HEADERS = ["Date", "Machine", "Tokens", "Cost ($)"]

# Column headers for machine totals sheet
MACHINE_HEADERS = ["Machine", "Tokens", "Cost ($)"]

# Totals row label
TOTALS_LABEL = "TOTALS"


def get_machine_id():
    """Get a deterministic machine identifier based on hostname + MAC address."""
    hostname = platform.node()
    mac = uuid.getnode()

    machine_str = f"{hostname}-{mac}"
    hash_bytes = hashlib.sha256(machine_str.encode()).hexdigest()[:8]

    return f"{hostname}-{hash_bytes}"


def fetch_usage_data(period):
    """Fetch usage data from bunx ccusage for the given period."""
    try:
        print(f"Fetching {period} data...")
        result = subprocess.run(
            ['bunx', 'ccusage', period, '--json'],
            capture_output=True,
            text=True,
            check=True,
            encoding='utf-8'
        )

        data = json.loads(result.stdout)
        entries = data.get(period, [])

        if entries:
            print(f"  [OK] {len(entries)} entries")
        else:
            print(f"  [X] No entries found")

        return entries

    except subprocess.CalledProcessError as e:
        print(f"  [X] Error: {e.stderr}")
        return None
    except json.JSONDecodeError as e:
        print(f"  [X] JSON parse error: {e}")
        return None
    except Exception as e:
        print(f"  [X] Error: {e}")
        return None


def load_or_create_workbook():
    """Load existing workbook or create a new one."""
    excel_path = Path(EXCEL_FILE)

    if excel_path.exists():
        print(f"Loading: {EXCEL_FILE}")
        try:
            return load_workbook(EXCEL_FILE)
        except Exception as e:
            print(f"  [X] Error loading: {e}")
            print("  Creating new workbook...")

    print(f"Creating: {EXCEL_FILE}")
    wb = Workbook()

    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])

    # Create time-based sheets
    for sheet_name in ["Daily", "Weekly", "Monthly"]:
        ws = wb.create_sheet(sheet_name)
        setup_time_sheet_headers(ws)

    # Create machine totals sheet
    ws = wb.create_sheet("Machine Totals")
    setup_machine_sheet_headers(ws)

    return wb


def setup_time_sheet_headers(ws):
    """Set up headers for time-based sheets."""
    for col_idx, header in enumerate(TIME_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def setup_machine_sheet_headers(ws):
    """Set up headers for machine totals sheet."""
    for col_idx, header in enumerate(MACHINE_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def normalize_date_str(value):
    """
    Normalize any date value to a consistent string format.
    Handles: datetime objects, strings with time components, various formats.
    Returns the canonical string form for storage and comparison.
    """
    if value is None:
        return None

    # Handle datetime objects (openpyxl converts date strings to datetime)
    if hasattr(value, 'strftime'):
        return value.strftime('%Y-%m-%d')

    s = str(value).strip()

    # Remove time component if present (e.g., "2025-12-15 00:00:00")
    if ' ' in s:
        s = s.split(' ')[0]

    return s


def get_sort_key(date_str):
    """
    Convert a date string to a sortable tuple key.
    Handles:
      - Daily: "2025-12-15" -> (2025, 12, 15, 0)
      - Weekly: "2025-W02" or "2025-W2" -> (2025, 0, 2, 0)
      - Monthly: "2025-12" -> (2025, 12, 0, 0)
    """
    if date_str is None or date_str == TOTALS_LABEL:
        return (9999, 99, 99, 99)  # Sort totals last

    s = normalize_date_str(date_str)

    # Weekly format: 2025-W02 or 2025-W2
    week_match = re.match(r'^(\d{4})-W(\d{1,2})$', s)
    if week_match:
        year, week = int(week_match.group(1)), int(week_match.group(2))
        return (year, 0, week, 0)

    # Monthly format: 2025-12
    month_match = re.match(r'^(\d{4})-(\d{1,2})$', s)
    if month_match:
        year, month = int(month_match.group(1)), int(month_match.group(2))
        return (year, month, 0, 0)

    # Daily format: 2025-12-15
    day_match = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})$', s)
    if day_match:
        year, month, day = int(day_match.group(1)), int(day_match.group(2)), int(day_match.group(3))
        return (year, month, day, 0)

    # Fallback: return string for lexicographic sort
    return (0, 0, 0, s)


def find_existing_row(sheet, date_str, machine_id):
    """Find row with matching date and machine."""
    target = normalize_date_str(date_str)

    for row_num in range(2, sheet.max_row + 1):
        cell_date = sheet.cell(row=row_num, column=1).value
        cell_machine = sheet.cell(row=row_num, column=2).value

        if cell_date == TOTALS_LABEL:
            continue

        if normalize_date_str(cell_date) == target and cell_machine == machine_id:
            return row_num
    return None


def find_machine_row(sheet, machine_id):
    """Find row for a specific machine in Machine Totals sheet."""
    for row_num in range(2, sheet.max_row + 1):
        cell_machine = sheet.cell(row=row_num, column=1).value

        if cell_machine == TOTALS_LABEL:
            continue

        if cell_machine == machine_id:
            return row_num
    return None


def find_totals_row(sheet):
    """Find the TOTALS row."""
    for row_num in range(2, sheet.max_row + 1):
        if sheet.cell(row=row_num, column=1).value == TOTALS_LABEL:
            return row_num
    return None


def get_insert_position(sheet, date_str):
    """Find position to insert new row (sorted chronologically)."""
    totals_row = find_totals_row(sheet)
    max_row = totals_row if totals_row else sheet.max_row + 1
    target_key = get_sort_key(date_str)

    for row_num in range(2, max_row):
        cell_date = sheet.cell(row=row_num, column=1).value
        cell_key = get_sort_key(cell_date)
        if cell_key > target_key:
            return row_num

    return max_row


def update_sheet(sheet, entries, machine_id):
    """Update sheet with entries."""
    if not entries:
        return False

    rows_added = 0
    rows_updated = 0

    for entry in entries:
        date_str = entry.get('date') or entry.get('week') or entry.get('month', 'Unknown')

        existing_row = find_existing_row(sheet, date_str, machine_id)

        if existing_row:
            row_num = existing_row
            rows_updated += 1
        else:
            insert_pos = get_insert_position(sheet, date_str)
            sheet.insert_rows(insert_pos)
            row_num = insert_pos
            rows_added += 1

        # Store date as string (prevents Excel auto-conversion issues)
        sheet.cell(row=row_num, column=1, value=str(date_str))
        sheet.cell(row=row_num, column=2, value=machine_id)
        sheet.cell(row=row_num, column=3, value=entry.get('totalTokens', 0))
        sheet.cell(row=row_num, column=4, value=entry.get('totalCost', 0))

    update_totals_row(sheet)
    print(f"  Added {rows_added}, updated {rows_updated} rows")
    return True


def update_totals_row(sheet):
    """Calculate and update the totals row for time-based sheets."""
    totals_row = find_totals_row(sheet)

    if not totals_row:
        totals_row = sheet.max_row + 1

    total_tokens = 0
    total_cost = 0.0

    for row_num in range(2, sheet.max_row + 1):
        if sheet.cell(row=row_num, column=1).value == TOTALS_LABEL:
            continue

        total_tokens += sheet.cell(row=row_num, column=3).value or 0
        total_cost += sheet.cell(row=row_num, column=4).value or 0

    sheet.cell(row=totals_row, column=1, value=TOTALS_LABEL)
    sheet.cell(row=totals_row, column=2, value="")
    sheet.cell(row=totals_row, column=3, value=total_tokens)
    sheet.cell(row=totals_row, column=4, value=total_cost)


def update_machine_totals(wb):
    """Update the Machine Totals sheet by aggregating from Daily sheet."""
    if "Machine Totals" not in wb.sheetnames:
        ws = wb.create_sheet("Machine Totals")
        setup_machine_sheet_headers(ws)

    machine_sheet = wb["Machine Totals"]

    # Clear existing data (keep headers)
    if machine_sheet.max_row > 1:
        machine_sheet.delete_rows(2, machine_sheet.max_row)

    # Aggregate from Daily sheet
    if "Daily" not in wb.sheetnames:
        return

    daily_sheet = wb["Daily"]
    machine_totals = {}

    for row_num in range(2, daily_sheet.max_row + 1):
        date_val = daily_sheet.cell(row=row_num, column=1).value
        if date_val == TOTALS_LABEL:
            continue

        machine = daily_sheet.cell(row=row_num, column=2).value
        tokens = daily_sheet.cell(row=row_num, column=3).value or 0
        cost = daily_sheet.cell(row=row_num, column=4).value or 0

        if machine not in machine_totals:
            machine_totals[machine] = {'tokens': 0, 'cost': 0.0}

        machine_totals[machine]['tokens'] += tokens
        machine_totals[machine]['cost'] += cost

    # Write machine totals
    row_num = 2
    grand_tokens = 0
    grand_cost = 0.0

    for machine in sorted(machine_totals.keys()):
        data = machine_totals[machine]
        machine_sheet.cell(row=row_num, column=1, value=machine)
        machine_sheet.cell(row=row_num, column=2, value=data['tokens'])
        machine_sheet.cell(row=row_num, column=3, value=data['cost'])
        grand_tokens += data['tokens']
        grand_cost += data['cost']
        row_num += 1

    # Write grand totals
    machine_sheet.cell(row=row_num, column=1, value=TOTALS_LABEL)
    machine_sheet.cell(row=row_num, column=2, value=grand_tokens)
    machine_sheet.cell(row=row_num, column=3, value=grand_cost)

    print(f"Machine Totals sheet:")
    print(f"  {len(machine_totals)} machines")


def format_time_sheet(sheet):
    """Apply formatting to time-based sheets."""
    widths = [12, 28, 15, 12]
    for idx, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    for row_num in range(2, sheet.max_row + 1):
        date_val = sheet.cell(row=row_num, column=1).value

        # Right-align numeric columns
        sheet.cell(row=row_num, column=3).alignment = Alignment(horizontal="right")
        sheet.cell(row=row_num, column=4).alignment = Alignment(horizontal="right")

        # Number formats
        sheet.cell(row=row_num, column=3).number_format = '#,##0'
        sheet.cell(row=row_num, column=4).number_format = '$#,##0.00'

        # Style TOTALS row
        if date_val == TOTALS_LABEL:
            for col_idx in range(1, 5):
                cell = sheet.cell(row=row_num, column=col_idx)
                cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
                cell.font = Font(bold=True)
                cell.border = Border(top=Side(style='thin'))


def format_machine_sheet(sheet):
    """Apply formatting to machine totals sheet."""
    widths = [28, 15, 12]
    for idx, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    for row_num in range(2, sheet.max_row + 1):
        machine_val = sheet.cell(row=row_num, column=1).value

        # Right-align numeric columns
        sheet.cell(row=row_num, column=2).alignment = Alignment(horizontal="right")
        sheet.cell(row=row_num, column=3).alignment = Alignment(horizontal="right")

        # Number formats
        sheet.cell(row=row_num, column=2).number_format = '#,##0'
        sheet.cell(row=row_num, column=3).number_format = '$#,##0.00'

        # Style TOTALS row
        if machine_val == TOTALS_LABEL:
            for col_idx in range(1, 4):
                cell = sheet.cell(row=row_num, column=col_idx)
                cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
                cell.font = Font(bold=True)
                cell.border = Border(top=Side(style='thin'))


def display_summary(wb):
    """Display summary."""
    print("\n" + "=" * 45)
    print("SUMMARY")
    print("=" * 45)

    for sheet_name in ["Daily", "Weekly", "Monthly"]:
        if sheet_name not in wb.sheetnames:
            continue

        sheet = wb[sheet_name]
        totals_row = find_totals_row(sheet)

        if totals_row:
            tokens = sheet.cell(row=totals_row, column=3).value or 0
            cost = sheet.cell(row=totals_row, column=4).value or 0
            entry_count = totals_row - 2

            print(f"\n{sheet_name}: {entry_count} entries")
            print(f"  Tokens: {tokens:,}")
            print(f"  Cost:   ${cost:.2f}")

    # Machine totals summary
    if "Machine Totals" in wb.sheetnames:
        sheet = wb["Machine Totals"]
        print(f"\nMachine Totals:")
        for row_num in range(2, sheet.max_row + 1):
            machine = sheet.cell(row=row_num, column=1).value
            tokens = sheet.cell(row=row_num, column=2).value or 0
            cost = sheet.cell(row=row_num, column=3).value or 0
            if machine == TOTALS_LABEL:
                print(f"  {machine}: {tokens:,} tokens, ${cost:.2f}")
            else:
                print(f"  {machine}: ${cost:.2f}")

    print("\n" + "=" * 45)


def main():
    """Main execution flow."""
    print("\n" + "=" * 45)
    print("AI Usage Budget Tracker")
    print("=" * 45)

    machine_id = get_machine_id()
    print(f"Machine: {machine_id}\n")

    # Fetch data
    daily_entries = fetch_usage_data('daily')
    weekly_entries = fetch_usage_data('weekly')
    monthly_entries = fetch_usage_data('monthly')

    if not daily_entries and not weekly_entries and not monthly_entries:
        print("\n[X] No data fetched. Check bunx ccusage.")
        sys.exit(1)

    # Update workbook
    print("\nUpdating Excel...")
    wb = load_or_create_workbook()

    if daily_entries:
        print("Daily sheet:")
        if "Daily" not in wb.sheetnames:
            ws = wb.create_sheet("Daily")
            setup_time_sheet_headers(ws)
        update_sheet(wb["Daily"], daily_entries, machine_id)
        format_time_sheet(wb["Daily"])

    if weekly_entries:
        print("Weekly sheet:")
        if "Weekly" not in wb.sheetnames:
            ws = wb.create_sheet("Weekly")
            setup_time_sheet_headers(ws)
        update_sheet(wb["Weekly"], weekly_entries, machine_id)
        format_time_sheet(wb["Weekly"])

    if monthly_entries:
        print("Monthly sheet:")
        if "Monthly" not in wb.sheetnames:
            ws = wb.create_sheet("Monthly")
            setup_time_sheet_headers(ws)
        update_sheet(wb["Monthly"], monthly_entries, machine_id)
        format_time_sheet(wb["Monthly"])

    # Update machine totals (aggregates from Daily)
    update_machine_totals(wb)
    format_machine_sheet(wb["Machine Totals"])

    # Save
    try:
        wb.save(EXCEL_FILE)
        print(f"\n[OK] Saved to {EXCEL_FILE}")
        display_summary(wb)
    except PermissionError:
        print(f"\n[X] Cannot save - {EXCEL_FILE} is open in another program.")
        print("    Please close Excel and run again.")
        sys.exit(1)
    except Exception as e:
        print(f"\n[X] Error saving: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
